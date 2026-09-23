"""Test harness.

Real components: PostgreSQL (schema built by Alembic migrations), ChromaDB
(persistent, temp directory), the FastAPI app, BM25 and the ingestion pipeline.

Faked (no network / no model downloads in CI): the sentence-transformer embedding
model (a deterministic hashed bag-of-words embedder) and Gemini (a scriptable
fake).  Both are injected through the same seams the app uses in production.
"""

from __future__ import annotations

import hashlib
import io
import json
import math
import os
import re
import shutil
import tempfile
from pathlib import Path

_TMP = tempfile.mkdtemp(prefix="edurag-tests-")
os.environ.update(
    {
        "DATABASE_URL": os.environ.get(
            "TEST_DATABASE_URL", "postgresql://postgres:postgres@localhost:5432/edurag_test"
        ),
        "SECRET_KEY": "test-secret-key-which-is-long-enough-0123456789",
        "ENVIRONMENT": "test",
        "VECTOR_DB_PATH": f"{_TMP}/vectors",
        "UPLOAD_DIR": f"{_TMP}/uploads",
        "GEMINI_API_KEY": "AIzaSyFAKE-KEY-FOR-TESTS-0123456789abcdefgh",
        "PRELOAD_MODELS": "false",
        "RUN_STARTUP_RECOVERY": "false",
        "RERANKER_ENABLED": "false",
        "CORS_ORIGINS": "*",
    }
)

import pytest  # noqa: E402
from alembic import command  # noqa: E402
from alembic.config import Config  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from langchain_core.embeddings import Embeddings  # noqa: E402
from sqlalchemy import text  # noqa: E402

BACKEND_DIR = Path(__file__).resolve().parents[1]


# ------------------------------------------------------------------- fake embeddings
class HashEmbeddings(Embeddings):
    """Deterministic bag-of-words embedder: similar wording -> higher cosine."""

    dim = 384

    def _embed(self, text_value: str) -> list[float]:
        vector = [0.0] * self.dim
        # Like a real semantic model, ignore function words; never return a zero vector.
        for token in (tokenize(text_value) or ["__empty__"]):
            bucket = int(hashlib.md5(token.encode()).hexdigest(), 16) % self.dim
            vector[bucket] += 1.0
        norm = math.sqrt(sum(v * v for v in vector)) or 1.0
        return [v / norm for v in vector]

    def embed_documents(self, texts):
        return [self._embed(t) for t in texts]

    def embed_query(self, text_value):
        return self._embed(text_value)


import app.ai.vectorstore.chroma_service as chroma_module  # noqa: E402
from app.ai.keyword.bm25_index import tokenize  # noqa: E402

_EMBEDDINGS = HashEmbeddings()
chroma_module.get_embeddings = lambda: _EMBEDDINGS  # type: ignore[assignment]


# ----------------------------------------------------------------------- fake Gemini
NOTES_JSON = json.dumps(
    [
        {"title": "Photosynthesis", "content": "Plants convert **light** into chemical energy.\n\n- Chlorophyll absorbs light\n- Produces glucose", "keywords": ["chlorophyll", "glucose", "light"]},
        {"title": "Respiration", "content": "Mitochondria produce ATP through cellular respiration.", "keywords": ["ATP", "mitochondria"]},
    ]
)
FLASHCARDS_JSON = json.dumps(
    [{"question": f"Question {i}?", "answer": f"Answer {i}."} for i in range(1, 6)]
)
QUIZ_JSON = json.dumps(
    [
        {"question": f"Quiz question {i}?", "option_a": f"A{i}", "option_b": f"B{i}", "option_c": f"C{i}", "option_d": f"D{i}", "correct_answer": "ABCD"[i % 4], "explanation": f"Because {i}."}
        for i in range(1, 6)
    ]
)


def plan_json(days: int = 7) -> str:
    return json.dumps(
        [
            {"day": d, "time": "09:00 AM", "title": f"Topic {d}", "description": f"Study topic {d} carefully.", "duration": "60 min", "quote": "Keep going."}
            for d in range(1, days + 1)
        ]
    )


class FakeLLM:
    """Scriptable stand-in for GeminiService."""

    def __init__(self) -> None:
        self.reset()

    def reset(self) -> None:
        self.script: list = []
        self.calls: list[tuple[str | None, str]] = []

    def queue(self, *responses) -> None:
        self.script.extend(responses)

    def _next(self, user: str, system: str | None) -> str:
        self.calls.append((system, user))
        if self.script:
            item = self.script.pop(0)
            if isinstance(item, Exception):
                raise item
            return item
        return self._default(user, system or "")

    def generate(self, user: str, system: str | None = None) -> str:
        return self._next(user, system)

    def stream(self, user: str, system: str | None = None):
        text_value = self._next(user, system)
        for word in re.findall(r"\S+\s*", text_value):
            yield word

    @staticmethod
    def _default(user: str, system: str) -> str:
        if "revision notes" in system:
            return NOTES_JSON
        if "flashcards" in system:
            return FLASHCARDS_JSON
        if "multiple-choice" in system:
            return QUIZ_JSON
        if "study schedules" in system:
            days = re.search(r"Create a (\d+)-day", system)
            return plan_json(int(days.group(1)) if days else 7)
        if "rewrite follow-up" in system:
            return re.search(r"Current question: (.*)", user).group(1)
        # teacher: cite the first source
        match = re.search(r'<source id="1"[^>]*>\n(.*?)\n</source>', user, re.DOTALL)
        snippet = (match.group(1)[:80] if match else "nothing").replace("\n", " ")
        return f"According to the material [1]: {snippet}"


import app.ai.llm.gemini_service as gemini_module  # noqa: E402

FAKE_LLM = FakeLLM()
gemini_module.get_llm = lambda: FAKE_LLM  # type: ignore[assignment]

from app.core.config import settings  # noqa: E402
from app.core.rate_limit import limiter  # noqa: E402
from app.database.session import SessionLocal, engine  # noqa: E402
from app.main import app  # noqa: E402
import app.services.generation_base as generation_base  # noqa: E402
import app.services.ingestion_service as ingestion  # noqa: E402
from app.ai.vectorstore.chroma_service import get_vectorstore, reset_vectorstore  # noqa: E402
from app.ai.retrieval.retriever import reset_retriever  # noqa: E402


# --------------------------------------------------------------------------- fixtures
@pytest.fixture(scope="session", autouse=True)
def _database():
    with engine.begin() as conn:
        conn.execute(text("DROP SCHEMA public CASCADE"))
        conn.execute(text("CREATE SCHEMA public"))
    cfg = Config(str(BACKEND_DIR / "alembic.ini"))
    command.upgrade(cfg, "head")
    yield
    shutil.rmtree(_TMP, ignore_errors=True)


@pytest.fixture(autouse=True)
def _clean_state(monkeypatch):
    """Isolate every test: empty tables, empty vector store, fresh fakes."""
    with engine.begin() as conn:
        conn.execute(text("TRUNCATE users CASCADE"))
    store = get_vectorstore()
    existing = store._collection.get(include=[])["ids"]
    if existing:
        store._collection.delete(ids=existing)
    reset_retriever()
    limiter.reset()
    generation_base._in_flight.clear()
    FAKE_LLM.reset()
    shutil.rmtree(settings.UPLOAD_DIR, ignore_errors=True)

    # Process uploads inline so tests are deterministic (the async path has its own test).
    monkeypatch.setattr(ingestion.ingestion_queue, "submit", lambda document_id: ingestion.process_document(document_id))
    yield


@pytest.fixture
def llm() -> FakeLLM:
    return FAKE_LLM


@pytest.fixture
def vectors():
    return get_vectorstore()


@pytest.fixture
def db():
    session = SessionLocal()
    try:
        yield session
    finally:
        session.close()


@pytest.fixture
def client() -> TestClient:
    return TestClient(app, raise_server_exceptions=False)


class Api:
    """Small helper bundling a client with one user's auth header."""

    def __init__(self, client: TestClient, email: str, password: str = "correct-horse-1"):
        self.client = client
        self.email = email
        self.password = password
        r = client.post("/auth/register", json={"full_name": email.split("@")[0].title(), "email": email, "password": password})
        assert r.status_code == 201, r.text
        self.user_id = r.json()["id"]
        r = client.post("/auth/login", json={"email": email, "password": password})
        assert r.status_code == 200, r.text
        self.headers = {"Authorization": f"Bearer {r.json()['access_token']}"}

    def get(self, url, **kw):
        return self.client.get(url, headers=self.headers, **kw)

    def post(self, url, **kw):
        return self.client.post(url, headers=self.headers, **kw)

    def put(self, url, **kw):
        return self.client.put(url, headers=self.headers, **kw)

    def patch(self, url, **kw):
        return self.client.patch(url, headers=self.headers, **kw)

    def delete(self, url, **kw):
        return self.client.delete(url, headers=self.headers, **kw)

    # ---- convenience workflows
    def subject(self, name="Biology") -> str:
        r = self.post("/subjects", json={"name": name, "description": "test"})
        assert r.status_code == 201, r.text
        return r.json()["id"]

    def upload(self, subject_id, data: bytes, filename="notes.pdf", content_type="application/pdf"):
        return self.post(f"/documents/upload/{subject_id}", files={"file": (filename, data, content_type)})

    def upload_settled(self, subject_id, data: bytes, filename="notes.pdf", content_type="application/pdf") -> dict:
        """Upload, then return the document AFTER processing (tests process inline)."""
        r = self.upload(subject_id, data, filename, content_type)
        assert r.status_code == 200, r.text
        return self.get(f"/documents/{r.json()['id']}/status").json()

    def upload_pdf(self, subject_id, pages=None, filename="biology.pdf") -> dict:
        return self.upload_settled(subject_id, make_pdf(pages or DEFAULT_PAGES), filename)


@pytest.fixture
def alice(client) -> Api:
    return Api(client, "alice@example.com")


@pytest.fixture
def bob(client) -> Api:
    return Api(client, "bob@example.com")


# ------------------------------------------------------------------------ file makers
DEFAULT_PAGES = [
    "Photosynthesis is the process by which plants convert light energy into chemical energy. "
    "Chlorophyll in the chloroplast absorbs sunlight and produces glucose and oxygen from carbon dioxide and water. "
    "The light reactions occur in the thylakoid membranes while the Calvin cycle occurs in the stroma.",
    "Cellular respiration takes place in the mitochondria. Glucose is broken down to release ATP, the energy currency of the cell. "
    "Glycolysis, the Krebs cycle and the electron transport chain are the three main stages of respiration.",
    "DNA replication is semi-conservative. Helicase unwinds the double helix and DNA polymerase synthesises the new strands. "
    "Okazaki fragments are joined by ligase on the lagging strand.",
]


def make_pdf(pages: list[str]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    for page_text in pages:
        y = 800
        words = page_text.split()
        line = ""
        for word in words:
            if len(line) + len(word) > 85:
                pdf.drawString(40, y, line)
                y -= 16
                line = ""
            line += word + " "
        if line:
            pdf.drawString(40, y, line)
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def make_blank_pdf(pages: int = 2) -> bytes:
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    for _ in range(pages):
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def make_docx(paragraphs: list[str]) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("Chapter One", level=1)
    for p in paragraphs:
        doc.add_paragraph(p)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def make_pptx(slides: list[str]) -> bytes:
    from pptx import Presentation

    prs = Presentation()
    for body in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Slide"
        slide.placeholders[1].text = body
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
