from docx import Document as DocxDocument
from langchain_core.documents import Document

from app.ai.guardrails.sanitize import clean_text
from app.ai.loaders.base import DocumentProcessingError


class DOCXLoaderService:
    def load(self, path: str) -> list[Document]:
        try:
            doc = DocxDocument(path)
        except Exception as exc:
            raise DocumentProcessingError(
                "We couldn't open this Word document. It may be corrupted."
            ) from exc

        parts: list[str] = []
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "") if paragraph.style is not None else ""
            if style.lower().startswith("heading"):
                parts.append(f"\n{text}\n")
            else:
                parts.append(text)
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))

        text = clean_text("\n".join(parts))
        if not text:
            return []
        return [Document(page_content=text, metadata={"unit": "page"})]
