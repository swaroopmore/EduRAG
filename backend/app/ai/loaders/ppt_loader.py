from langchain_core.documents import Document
from pptx import Presentation

from app.ai.guardrails.sanitize import clean_text
from app.ai.loaders.base import DocumentProcessingError


class PPTLoaderService:
    """PPTX loader (python-pptx).  Each slide becomes one "page"."""

    def load(self, path: str) -> list[Document]:
        try:
            presentation = Presentation(path)
        except Exception as exc:
            raise DocumentProcessingError(
                "We couldn't open this presentation. It may be corrupted."
            ) from exc

        slides = list(presentation.slides)
        documents: list[Document] = []
        for index, slide in enumerate(slides):
            texts: list[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                    texts.append(shape.text_frame.text)
                if getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        texts.append(" | ".join(c.text for c in row.cells if c.text.strip()))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                texts.append(slide.notes_slide.notes_text_frame.text)
            text = clean_text("\n".join(t for t in texts if t and t.strip()))
            if text:
                documents.append(
                    Document(
                        page_content=text,
                        metadata={"page": index, "unit": "slide", "total_pages": len(slides)},
                    )
                )
        return documents
